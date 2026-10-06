"""Wrap a text file as a Word document whose body is exactly the file's text (one paragraph per line).
Spelling/grammar checks are switched off so Word doesn't stall on code on a slow machine."""
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_LINE_SPACING
src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Consolas"; st.font.size = Pt(7)
pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
st.element.get_or_add_rPr().append(OxmlElement("w:noProof"))
settings = doc.settings.element
for tag in ("w:hideSpellingErrors", "w:hideGrammaticalErrors"):
    settings.append(OxmlElement(tag))
body = doc.element.body
for ln in lines:
    p = doc.add_paragraph(); r = p.add_run(ln)
    t = r._r.find(qn("w:t"))
    if t is not None: t.set(qn("xml:space"), "preserve")
doc.save(out)
# verify the text comes back exactly
back = "\n".join(p.text for p in Document(out).paragraphs)
src_txt = "\n".join(lines)
assert back.rstrip("\n") == src_txt.rstrip("\n"), "round-trip mismatch"
print(out, "ok", len(lines), "paragraphs")
